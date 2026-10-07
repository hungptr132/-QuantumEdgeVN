"""One test per defect a code review found in the first statement extractor.

Each of these wrote, or could write, a wrong figure carrying clean provenance — the single
outcome `CLAUDE.md` says must not happen. Two of them contradicted claims made in the
commit message that shipped them, which is the part worth remembering: the extractor's own
docstring asserted that a statement must satisfy the VAS identities, and that was only
true when every input happened to be present.

They live in their own file because they share no cause, only a severity.
"""

from decimal import Decimal
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from quantumedge.config import load_settings
from quantumedge.db.enums import PipelineStage
from quantumedge.extract.mapping.vas import load_spec, spec_for
from quantumedge.extract.statements import extract_statements, interpret, read_document
from quantumedge.ingest.base import store_raw
from quantumedge.ingest.filings import archive, scan, to_payload
from quantumedge.runs import start_run


@pytest.fixture(scope="module")
def spec():
    return spec_for(load_settings().rules_dir)


@pytest.fixture
def filings_dir(tmp_path: Path) -> Path:
    d = tmp_path / "filings" / "TST"
    d.mkdir(parents=True)
    return d


def _ingest(session: Session, filings_dir: Path, store: Path) -> int:
    session.execute(
        text(
            "INSERT INTO source (code, kind, name) "
            "VALUES ('company_filing', 'manual', 'Filings') ON CONFLICT DO NOTHING"
        )
    )
    run_id = start_run(session, PipelineStage.INGEST, "t", source_code="company_filing")
    for found in scan(filings_dir.parent).found:
        stored = archive(found, store)
        store_raw(session, run_id, [to_payload(found, stored)])
    return run_id


def _transcription(path: Path, **overrides: object) -> Path:
    """A balanced hand transcription, as a person would fill the template in."""
    facts: dict[str, object] = {
        "total_assets": 1_000_000_000,
        "total_liabilities": 400_000_000,
        "total_equity": 600_000_000,
        "total_liabilities_and_equity": 1_000_000_000,
        "revenue": 800_000_000,
        "cost_of_goods_sold": 560_000_000,
        "gross_profit": 240_000_000,
        "net_profit_after_tax": 68_000_000,
    }
    facts.update(overrides)
    yaml.safe_dump(
        {
            "source": {"document": "BCTC 2024", "unit": "đồng", "transcribed_by": "richard"},
            "facts": {
                code: {"value": value, "cite": "Bang CDKT, trang 5"}
                for code, value in facts.items()
            },
        },
        path.open("w", encoding="utf-8"),
        allow_unicode=True,
    )
    return path


def _edit(path: Path, mutate: object) -> None:
    """Load, mutate, save — for tests that need a transcription slightly broken."""
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    mutate(data)  # type: ignore[operator]
    yaml.safe_dump(data, path.open("w", encoding="utf-8"), allow_unicode=True)


def _facts(session: Session, ticker_id: int) -> dict[str, Decimal]:
    return {
        str(code): value
        for code, value in session.execute(
            text(
                "SELECT ff.fact_code, ff.value FROM financial_fact ff "
                "JOIN financial_period fp ON fp.id = ff.financial_period_id "
                "WHERE fp.ticker_id = :t"
            ),
            {"t": ticker_id},
        ).all()
    }


def _by_code(spec, path: Path) -> dict:
    out = {}
    for row in read_document(path).rows:
        candidate = interpret(spec, row)
        if candidate is not None:
            out.setdefault(candidate.vas_code, candidate)
    return out


def _run(session: Session, spec) -> object:
    return extract_statements(session, start_run(session, PipelineStage.EXTRACT, "t"), spec)


class TestPeriodIdentity:
    def test_h1_and_h2_of_one_year_are_separate_periods(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        """Both halves mapped to period_type 'H' with fiscal_quarter NULL, and the unique
        key is COALESCE(fiscal_quarter, 0) — so 2024-H2 found 2024-H1's period row and
        upserted its figures over it, under a period_end of 30 June."""
        _transcription(filings_dir / "2024-H1-transcribed.yaml")
        _transcription(filings_dir / "2024-H2-transcribed.yaml", net_profit_after_tax=99_000_000)
        _ingest(session, filings_dir, tmp_path / "store")

        _run(session, spec)

        periods = session.execute(
            text(
                "SELECT fiscal_quarter, period_end FROM financial_period "
                "WHERE ticker_id = :t AND period_type = 'H' ORDER BY period_end"
            ),
            {"t": seed_ticker},
        ).all()
        assert len(periods) == 2, "H1 and H2 must not collide into one period"
        assert [str(p.period_end) for p in periods] == ["2024-06-30", "2024-12-31"]
        # For a half-year the column records the half, which is what makes the key unique.
        assert [p.fiscal_quarter for p in periods] == [1, 2]


class TestTheNewestDocumentWins:
    def test_a_later_filing_is_not_overwritten_by_an_earlier_one(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        """Every stored version was processed and each upserted over the last, so the final
        writer was the OLDEST: an audited filing ingested after an unaudited one was
        silently replaced by the superseded figures."""
        store = tmp_path / "store"

        # Ingested in two passes, because that is the real sequence: an unaudited filing
        # arrives, and the audited one weeks later. Doing both in one scan would order them
        # alphabetically, which is not what "newest" means.
        _transcription(filings_dir / "2024-Y-unaudited.yaml")
        _ingest(session, filings_dir, store)

        later = _transcription(filings_dir / "2024-Y-audited.yaml", net_profit_after_tax=99_000_000)

        def mark(data: dict) -> None:
            data["source"]["document"] = "BCTC da kiem toan"

        _edit(later, mark)
        _ingest(session, filings_dir, store)

        _run(session, spec)

        value, path = session.execute(
            text(
                "SELECT ff.value, ff.source_field_path FROM financial_fact ff "
                "JOIN financial_period fp ON fp.id = ff.financial_period_id "
                "WHERE fp.ticker_id = :t AND ff.fact_code = 'net_profit_after_tax'"
            ),
            {"t": seed_ticker},
        ).one()
        assert value == Decimal("99000000.0000"), "the later filing's figure must win"
        assert "kiem toan" in path


class TestZeroChecksIsNotAPass:
    def test_a_statement_with_no_testable_identity_is_refused(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        """The worst finding, because the docstring claimed the opposite. check_identities
        skips an identity whose inputs are absent, so a filing missing its totals produced
        zero checks, zero failures, and every figure written unverified — and the template
        actively invites it: "delete any line you cannot find"."""
        path = _transcription(filings_dir / "2024-Y-transcribed.yaml")

        def drop_totals(data: dict) -> None:
            for key in ("total_liabilities_and_equity", "gross_profit"):
                del data["facts"][key]

        _edit(path, drop_totals)
        _ingest(session, filings_dir, tmp_path / "store")

        outcome = _run(session, spec)

        assert outcome.facts_written == 0  # type: ignore[attr-defined]
        assert any("could be tested" in f for f in outcome.identity_failures)  # type: ignore[attr-defined]
        assert _facts(session, seed_ticker) == {}, "zero checks is not a pass"


class TestColumnConfusion:
    def test_a_note_reference_is_not_read_as_the_figure(self, filings_dir: Path, spec) -> None:
        """Circular 200's balance sheet runs Chỉ tiêu | Mã số | Thuyết minh | Số cuối năm,
        and many filings number the note column with bare integers. The first number after
        the code was therefore a note reference — and for the facts no identity covers, the
        statement still balanced."""
        path = filings_dir / "2024-Y-withnotes.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "CDKT"
        ws.append(["TÀI SẢN", "Mã số", "Thuyết minh", "Số cuối năm"])
        ws.append(["TỔNG CỘNG TÀI SẢN", "270", 24, 1_000_000_000])
        wb.save(path)

        candidates = _by_code(spec, path)

        assert candidates["270"].amount == Decimal(1_000_000_000), (
            "the note reference 24 must not be read as total assets"
        )

    def test_a_mapped_code_is_preferred_over_a_bare_row_number(
        self, filings_dir: Path, spec
    ) -> None:
        """A workbook numbering its own items bound the code to that number and never saw
        the real Mã số, losing the line silently."""
        path = filings_dir / "2024-Y-numbered.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "KQHDKD"
        # "1" is the item number; "10" is the Mã số for net revenue.
        ws.append(["1", "Doanh thu thuần về bán hàng", "10", 800_000_000])
        wb.save(path)

        candidates = _by_code(spec, path)

        assert "10" in candidates, "the mapped code must win over a bare row number"
        assert candidates["10"].amount == Decimal(800_000_000)


class TestDenomination:
    """The one error the identities provably cannot catch: all three are scale-invariant,
    so a statement in triệu đồng balances perfectly and is stored a million times small."""

    def test_a_statement_in_millions_is_refused(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        path = filings_dir / "2024-Y-millions.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "CDKT"
        ws.append(["BẢNG CÂN ĐỐI KẾ TOÁN", None, None])
        ws.append(["Đơn vị tính: triệu đồng", None, None])
        ws.append(["TỔNG CỘNG TÀI SẢN", "270", 1_000_000])
        ws.append(["TỔNG CỘNG NGUỒN VỐN", "440", 1_000_000])
        wb.save(path)
        _ingest(session, filings_dir, tmp_path / "store")

        outcome = _run(session, spec)

        assert outcome.facts_written == 0  # type: ignore[attr-defined]
        assert any("triệu đồng" in n for n in outcome.notes)  # type: ignore[attr-defined]

    def test_a_statement_declaring_dong_is_accepted(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        """The guard must not reject the normal case, which declares đồng explicitly."""
        path = filings_dir / "2024-Y-dong.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "CDKT"
        ws.append(["Đơn vị tính: đồng", None, None])
        ws.append(["TỔNG CỘNG TÀI SẢN", "270", 1_000_000_000])
        ws.append(["TỔNG CỘNG NGUỒN VỐN", "440", 1_000_000_000])
        wb.save(path)
        _ingest(session, filings_dir, tmp_path / "store")

        outcome = _run(session, spec)

        assert outcome.facts_written > 0  # type: ignore[attr-defined]

    def test_a_transcription_declaring_millions_is_refused(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        """Same trap by hand: the template says `value:` with no unit, so a transcriber
        copying from a triệu đồng column produces the same silent error."""
        path = _transcription(filings_dir / "2024-Y-transcribed.yaml")

        def set_unit(data: dict) -> None:
            data["source"]["unit"] = "triệu đồng"

        _edit(path, set_unit)
        _ingest(session, filings_dir, tmp_path / "store")

        outcome = _run(session, spec)

        assert outcome.facts_written == 0  # type: ignore[attr-defined]
        assert any("đồng" in n for n in outcome.notes)  # type: ignore[attr-defined]


class TestOneBadDocumentDoesNotStopTheRun:
    def test_an_unknown_transcribed_fact_code_refuses_only_that_document(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        """A typo violated the fact_code foreign key at insert time, outside the
        per-document handler — so one bad file aborted the run for every other symbol."""
        bad = _transcription(filings_dir / "2024-Q1-transcribed.yaml")

        def typo(data: dict) -> None:
            data["facts"]["total_equiy"] = data["facts"].pop("total_equity")

        _edit(bad, typo)
        # A good file for another period, which must survive.
        _transcription(filings_dir / "2024-Y-transcribed.yaml")
        _ingest(session, filings_dir, tmp_path / "store")

        outcome = _run(session, spec)

        assert any("not a known fact code" in n for n in outcome.notes)  # type: ignore[attr-defined]
        assert outcome.facts_written > 0, "the good document must still be extracted"  # type: ignore[attr-defined]

    def test_an_unreadable_figure_is_reported_not_dropped(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        """A figure typed with comma grouping vanished silently — and losing a total
        disabled the identity that would have caught the rest."""
        path = _transcription(filings_dir / "2024-Y-transcribed.yaml")

        def comma_grouped(data: dict) -> None:
            data["facts"]["total_assets"]["value"] = "1,000,000,000"

        _edit(path, comma_grouped)
        _ingest(session, filings_dir, tmp_path / "store")

        outcome = _run(session, spec)

        assert outcome.facts_written == 0  # type: ignore[attr-defined]
        assert any("could not read" in n for n in outcome.notes)  # type: ignore[attr-defined]


class TestDeductionSigns:
    def test_a_parenthesised_cost_does_not_void_the_document(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        """Some filings print Giá vốn hàng bán as (560.000.000). The gross-profit identity
        subtracts it, so a negative made the left side revenue PLUS cost, the identity
        failed, and a perfectly good balance sheet in the same workbook went with it."""
        path = filings_dir / "2024-Y-parens.xlsx"
        wb = Workbook()
        bs = wb.active
        bs.title = "CDKT"
        for row in [
            ["TỔNG CỘNG TÀI SẢN", "270", 1_000_000_000],
            ["NỢ PHẢI TRẢ", "300", 400_000_000],
            ["VỐN CHỦ SỞ HỮU", "400", 600_000_000],
            ["TỔNG CỘNG NGUỒN VỐN", "440", 1_000_000_000],
        ]:
            bs.append(row)
        inc = wb.create_sheet("KQHDKD")
        for row in [
            ["Doanh thu thuần về bán hàng", "10", "800.000.000"],
            ["Giá vốn hàng bán", "11", "(560.000.000)"],
            ["Lợi nhuận gộp", "20", "240.000.000"],
        ]:
            inc.append(row)
        wb.save(path)
        _ingest(session, filings_dir, tmp_path / "store")

        outcome = _run(session, spec)

        assert outcome.identity_failures == []  # type: ignore[attr-defined]
        facts = _facts(session, seed_ticker)
        # Stored as a magnitude: the identity that uses it already subtracts.
        assert facts["cost_of_goods_sold"] == Decimal("560000000.0000")


class TestIdempotencyIsReportable:
    def test_a_second_run_reports_unchanged_rather_than_rewritten(
        self,
        session: Session,
        seed_ticker: int,
        fact_defs: int,
        filings_dir: Path,
        tmp_path: Path,
        spec,
    ) -> None:
        """Invariant 7 says a second run reports `unchanged`. The counters incremented on
        every upsert regardless, so a re-run looked exactly like the first and the claim
        could not be checked from the output."""
        _transcription(filings_dir / "2024-Y-transcribed.yaml")
        _ingest(session, filings_dir, tmp_path / "store")

        first = _run(session, spec)
        second = _run(session, spec)

        assert first.facts_written > 0  # type: ignore[attr-defined]
        assert first.periods_written == 1  # type: ignore[attr-defined]
        assert second.facts_written == 0, "nothing changed, so nothing is written"  # type: ignore[attr-defined]
        assert second.facts_unchanged == first.facts_written  # type: ignore[attr-defined]
        assert second.periods_written == 0  # type: ignore[attr-defined]


class TestTheSpecRefusesAmbiguity:
    def test_a_duplicate_vas_code_is_rejected_at_load(self, tmp_path: Path) -> None:
        """Last-write-wins dropped a mapping silently. The registry is expected to grow,
        with codes that differ between revisions of the circular, so a duplicate is a
        likely mistake rather than an intention."""
        path = tmp_path / "dupes.yaml"
        path.write_text(
            "identities: []\n"
            "lines:\n"
            '  - {vas_code: "300", fact_code: total_liabilities, statement: balance,\n'
            '     label_keywords: ["no phai tra"], status: assumed}\n'
            '  - {vas_code: "300", fact_code: total_equity, statement: balance,\n'
            '     label_keywords: ["von chu"], status: assumed}\n',
            encoding="utf-8",
        )

        with pytest.raises(ValueError, match="duplicate vas_code"):
            load_spec(path)
